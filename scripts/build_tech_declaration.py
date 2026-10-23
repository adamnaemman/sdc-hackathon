"""
Fill the official SDC 2026 Technical & Tools Declaration template for overinfinity.

The template is copied first and the answers are written into its own tables, so
the original design and structure survive.

Run:  python build_tech_declaration.py
"""

import os
import re
import shutil

from docx import Document

TEMPLATE = r"C:\Users\AHMAD ADAM HAKIMI\Downloads\2. SDC Hackathon_ Technical Tools Declaration Template.docx"
OUTDIR = r"D:\sdc-hackathon\docs"
OUTFILE = os.path.join(OUTDIR, "SDC2026_overinfinity_TechDeclaration.docx")

# ---------------------------------------------------------------------------
# Content
# ---------------------------------------------------------------------------

LIVE_URL = "https://aura-workforce.replit.app/"
REPO_URL = "https://github.com/adamnaemman/sdc-hackathon"

TEAM = {
    "Team Name": "overinfinity",
    "Project / Solution Name": "AURA-Workforce",
    "Name (Leader)": "Ahmad Adam Hakimi bin Naeman",
    "Email & Phone (Leader)": "0damkim254@gmail.com / +60 18-344 5135",
    "University Affiliation(s)": "International Islamic University Malaysia (IIUM)",
    "Live Prototype URL": LIVE_URL,
    "Code Repository URL (public / access granted to judges)": REPO_URL,
}

ARCHITECTURE = (
    "AURA-Workforce is a two-service deployment on Replit: a React single-page app in the browser and an "
    "Express API mounted at /api. The API serves three things. First, a curated catalogue of five "
    "future-ready roles and their learning steps, each step pointing at a free, verified provider page. "
    "Second, a live trend layer: it fetches recent public job postings, normalises the text, counts how "
    "often each role's required skills appear, and returns a demand share stamped with source, fetch time "
    "and sample size; results are cached for six hours. Third, two AI endpoints on DeepSeek: one turns "
    "free text into a structured profile, and one receives the whole profile plus the catalogue and "
    "returns a grounded recommendation with an ordered pathway. Every role ID and module ID returned by "
    "the model is validated against the catalogue before it reaches the browser. No personal data is "
    "stored anywhere; the profile lives in browser memory only."
)

LIBRARIES = [
    ("React 19.1", "User interface", "MIT"),
    ("Vite 7", "Build tool and dev server", "MIT"),
    ("Tailwind CSS 4", "Styling system", "MIT"),
    ("Radix UI primitives", "Accessible UI components", "MIT"),
    ("TanStack Query 5", "Client data fetching and cache", "MIT"),
    ("Express 5", "HTTP API server", "MIT"),
    ("Zod 3", "Request and response validation", "MIT"),
    ("Orval 8", "OpenAPI to typed client generation", "MIT"),
    ("TypeScript 5.9", "Language and type checking", "Apache-2.0"),
    ("pnpm 10", "Workspace package manager", "MIT"),
    ("Lucide", "Icon set", "ISC"),
]

EXTERNAL_APIS = [
    ("DeepSeek API (api.deepseek.com)", "Language model for profile extraction and the AI planner", "Paid (team-funded)"),
    ("Arbeitnow Job Board API", "Live source of public job postings for the trend layer", "Free, no key"),
]

CLOUD = [
    ("Replit Core (deployment + secrets)", "Hosts the web app and the API; stores API keys as secrets", "Sponsor-provided, Singapore region"),
    ("GitHub", "Source control and public repository for judges", "Free"),
    ("RunPod (reserved)", "GPU credit reserved for Phase 2 model hosting; not used in this prototype", "Sponsor-provided, USD 20/team"),
]

MODELS = [
    ("deepseek-chat", "DeepSeek", "Pre-trained (no fine-tuning)", "Commercial API terms"),
]

DATASETS = [
    ("U.S. BLS Occupational Outlook Handbook (2025-2035 projections)", "https://www.bls.gov/ooh/", "Public", "Public domain (US Government)", "N"),
    ("WEF Future of Jobs Report 2025", "https://www.weforum.org/publications/the-future-of-jobs-report-2025/", "Public", "Quoted with attribution", "N"),
    ("Arbeitnow Job Board API (live postings)", "https://www.arbeitnow.com/api/job-board-api", "Public", "Public API, free, no key", "N"),
    ("AURA role and learning-step catalogue", "https://github.com/adamnaemman/sdc-hackathon", "Synthetic / curated by the team", "Team-owned, sources cited in app", "N"),
    ("Hays Asia Salary Guide 2025 (two figures, quoted)", "Cited in PRD section 2.1", "Public", "Quoted with attribution", "N"),
]

AI_TOOLS = [
    ("DeepSeek (chat + coding agent)", "Code generation, refactoring, debugging, documentation", "Major"),
    ("Replit Agent", "Answered Git workflow questions only; no code committed from it", "Minor"),
]

REVIEW_STATEMENT = (
    "Every AI-generated change was reviewed as a diff before running, then verified by a workspace-wide "
    "TypeScript typecheck and a successful build in Replit. Endpoints were tested with scripted requests, "
    "including a prompt-injection attempt that the model correctly ignored. Generated copy was checked "
    "against the catalogue for invented roles, courses and statistics, and two data bugs were found and "
    "fixed this way. The team can explain every file, and the scoring engine remains deterministic."
)

SPONSORS = [
    ("Replit", "Replit Core, one month per team member", "Hosting, deployment and secret storage for the prototype"),
    ("RunPod", "USD 20 GPU credit per team", "Not used in the preliminary prototype; reserved for Phase 2 model hosting"),
]

SPONSOR_ESSENTIAL = (
    "Replit Core is essential: the prototype is built, run and deployed there, and it hosts both the API "
    "and the web app. The RunPod credit is not used yet; the AI client is OpenAI-compatible, so moving the "
    "model to a self-hosted GPU on RunPod in Phase 2 is a configuration change rather than a rewrite. The "
    "DeepSeek API is paid by the team, not sponsor-provided."
)

DATA_HANDLING = (
    "The application collects no account data and stores nothing on the server. A profile exists only in "
    "browser memory and is discarded when the tab closes. Two optional actions leave the browser, and only "
    "when the user presses a button: a free-text work description (sent to our API, forwarded to DeepSeek, "
    "neither logged nor stored) and the structured profile summary sent with the AI plan request. That "
    "request contains skills, years of experience and target sectors, never a name or contact detail. There "
    "is no CV upload and no database. Server logs record method, URL, status and response time only."
)

LIMITATIONS = (
    "Hardest risk: the model inventing a job, course or statistic. Mitigation: it may only choose from a "
    "five-role, twenty-step catalogue, every returned ID is validated server-side, and its fit score is "
    "bounded and displayed next to the deterministic estimate. Second risk: the live job feed is a global, "
    "tech-oriented board, so it measures the direction of skill demand rather than Malaysian vacancy "
    "counts; the interface states this and shows the sample size. Third: five roles cannot serve every "
    "occupation, so the plan is required to say when no role is a strong match. Fourth: a single AI "
    "provider; the client is OpenAI-compatible, so a fallback endpoint is a configuration change."
)


def word_count(text: str) -> int:
    return len(re.findall(r"[A-Za-z0-9][A-Za-z0-9'\-]*", text))


def fill_cell(cell, lines) -> None:
    cell.text = ""
    first = cell.paragraphs[0]
    for index, line in enumerate(lines):
        target = first if index == 0 else cell.add_paragraph()
        target.add_run(line)


def add_rows(table, rows) -> None:
    """Fill the template's existing blank rows first, then add only what is missing.

    The template ships with a fixed number of empty rows (e.g. 6 for libraries).
    Appending with add_row() would leave those blanks above the data, so we
    reuse them and delete any surplus.
    """
    data_rows = list(table.rows)[1:]  # row 0 is the header

    for index, row in enumerate(data_rows):
        if index < len(rows):
            for column, value in enumerate(rows[index]):
                row.cells[column].text = value
        else:
            row._element.getparent().remove(row._element)

    for values in rows[len(data_rows):]:
        cells = table.add_row().cells
        for column, value in enumerate(values):
            cells[column].text = value


def main() -> None:
    os.makedirs(OUTDIR, exist_ok=True)

    outfile = OUTFILE
    try:
        shutil.copyfile(TEMPLATE, outfile)
    except PermissionError:
        outfile = OUTFILE.replace(".docx", "_v2.docx")
        shutil.copyfile(TEMPLATE, outfile)
        print("NOTE: %s is locked. Writing %s instead." % (OUTFILE, outfile))

    document = Document(outfile)
    tables = document.tables

    # 1. Team overview (table 1): rows are label/value pairs.
    for row in tables[1].rows:
        label = row.cells[0].text.strip()
        if label in TEAM:
            row.cells[1].text = TEAM[label]

    # 2.1 architecture
    fill_cell(tables[2].cell(0, 0), [ARCHITECTURE])

    # 2.2 libraries (header row already present)
    add_rows(tables[3], LIBRARIES)

    # 3.1 external APIs
    add_rows(tables[4], EXTERNAL_APIS)

    # 3.2 cloud platforms
    add_rows(tables[5], CLOUD)

    # 4.1 models
    add_rows(tables[6], MODELS)

    # 4.2 datasets
    add_rows(tables[7], DATASETS)

    # 4.3 AI-assisted development tools
    add_rows(tables[8], AI_TOOLS)
    fill_cell(tables[9].cell(0, 0), [REVIEW_STATEMENT])

    # 5 sponsor resources
    add_rows(tables[10], SPONSORS)
    fill_cell(tables[11].cell(0, 0), [SPONSOR_ESSENTIAL])

    # 6.1 / 6.2
    fill_cell(tables[12].cell(0, 0), [DATA_HANDLING])
    fill_cell(tables[13].cell(0, 0), [LIMITATIONS])

    # 7 checklist: tick every box (all statements are true for this submission).
    for row in tables[14].rows:
        row.cells[0].text = "\u2611"

    # Yes/No paragraphs: 0 -> generative AI used (Yes), 1 -> sponsor resource used (Yes)
    yes_no = [p for p in document.paragraphs if p.text.strip().startswith("\u2610 Yes")]
    for paragraph in yes_no[:2]:
        paragraph.text = "\u2611 Yes   \u2610 No"

    # Header deadline + declaration block
    for paragraph in document.paragraphs:
        stripped = paragraph.text.strip()
        if stripped == "Submission Deadline: [Insert date]":
            paragraph.text = "Submission Deadline: 25 October 2026"
        elif stripped == "[Signature]":
            paragraph.text = "Ahmad Adam Hakimi bin Naeman (Team Leader, overinfinity)"
        elif stripped == "[Date]":
            paragraph.text = "23 October 2026"

    document.save(outfile)

    print("saved:", outfile)
    print("word counts -> architecture:", word_count(ARCHITECTURE), "(limit 200)")
    print("               review     :", word_count(REVIEW_STATEMENT), "(limit 100)")
    print("               sponsor    :", word_count(SPONSOR_ESSENTIAL), "(limit 100)")
    print("               data       :", word_count(DATA_HANDLING), "(limit 150)")
    print("               limits     :", word_count(LIMITATIONS), "(limit 150)")


if __name__ == "__main__":
    main()
